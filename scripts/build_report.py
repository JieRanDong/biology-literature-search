#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""Build the literature-search deliverable workbook: 文献清单.xlsx.

Two sheets:
  推荐文献      — the curated reading list, in reading order, with download status
  全部检索结果  — every de-duplicated record from merged.json

The curated list is read from a selection.json written during Stage 5. Its
`status` field must be one of the three literals below; the buckets are asserted
to partition the list so a typo'd status cannot silently vanish from the counts.

Usage:
  uv run scripts/build_report.py --selection tmp_litsearch/selection.json \
    --merged tmp_litsearch/merged.json --out tmp_litsearch/文献清单.xlsx

Payload fields consumed:
  selection.json  {topic, search_date, backends, items[{order,title,journal,doi,
                   year,status,local_file,note}]}
                  `local_file` is the path of the paper relative to the delivery
                  directory, i.e. `_paper/NN_<slug>.pdf` (Stage 5.2), or "" when
                  nothing was retrieved. It is printed verbatim into the 本地文件
                  column, so it must be a path the user can actually follow.
  merged.json     produced by merge_results.py — records[] with doi, pmid, pmcid,
                  title, year, journal, url, access, oa_status, found_in, cited_by
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required — run this via `uv run scripts/build_report.py`, "
             "which installs it from the PEP-723 header.")

STATUS_PDF = "已下载 PDF"
STATUS_TEXT = "已下载全文（非 PDF）"
STATUS_MISS = "未下载"
VALID_STATUS = {STATUS_PDF, STATUS_TEXT, STATUS_MISS}

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT = Font(color="0563C1", underline="single")
ITALIC = Font(italic=True, size=9, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

STATUS_FILL = {
    STATUS_PDF: PatternFill("solid", fgColor="E2EFDA"),
    STATUS_TEXT: PatternFill("solid", fgColor="FFF2CC"),
    STATUS_MISS: PatternFill("solid", fgColor="FCE4E4"),
}

SEL_COLS = ["序号", "文章名", "期刊名", "地址", "年份", "获取状态", "本地文件", "备注"]
SEL_WIDTHS = [6, 52, 24, 42, 7, 18, 38, 40]
ALL_COLS = ["文章名", "期刊名", "地址", "年份", "来源后端", "开放获取", "PMID"]
ALL_WIDTHS = [58, 30, 42, 7, 20, 12, 12]


def doi_url(doi):
    doi = (doi or "").strip()
    if not doi:
        return ""
    if doi.lower().startswith("http"):
        return doi
    return f"https://doi.org/{doi}"


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def put_link(ws, row, col, url):
    cell = ws.cell(row=row, column=col, value=url or None)
    if url:
        cell.hyperlink = url
        cell.font = LINK_FONT


def sheet_selection(wb, sel):
    ws = wb.active
    ws.title = "推荐文献"
    ws.append(SEL_COLS)
    for it in sel.get("items", []):
        status = it.get("status", "")
        if status not in VALID_STATUS:
            sys.exit(f"selection.json item {it.get('order')!r} has status "
                     f"{status!r}; must be one of {sorted(VALID_STATUS)}")
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=it.get("order", ""))
        ws.cell(row=r, column=2, value=it.get("title", ""))
        ws.cell(row=r, column=3, value=it.get("journal", ""))
        put_link(ws, r, 4, doi_url(it.get("doi")))
        ws.cell(row=r, column=5, value=it.get("year", ""))
        sc = ws.cell(row=r, column=6, value=status)
        sc.fill = STATUS_FILL[status]
        ws.cell(row=r, column=7, value=it.get("local_file") or "—")
        ws.cell(row=r, column=8, value=it.get("note", ""))
        for c in range(1, len(SEL_COLS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP if c in (2, 3, 8) else TOP
    for col, w in zip("ABCDEFGH", SEL_WIDTHS):
        ws.column_dimensions[col].width = w
    style_header(ws, len(SEL_COLS))

    note = ws.max_row + 2
    ws.cell(row=note, column=1,
            value=f"检索日期 {sel.get('search_date', '')}｜后端 {sel.get('backends', '')}"
                  f"｜主题：{sel.get('topic', '')}")
    ws.cell(row=note, column=1).font = ITALIC
    return ws


def sheet_all(wb, records):
    ws = wb.create_sheet("全部检索结果")
    ws.append(ALL_COLS)
    for rec in records:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=rec.get("title") or "")
        ws.cell(row=r, column=2, value=rec.get("journal") or "")
        put_link(ws, r, 3, doi_url(rec.get("doi")) or rec.get("url") or "")
        ws.cell(row=r, column=4, value=rec.get("year") or "")
        ws.cell(row=r, column=5, value=", ".join(rec.get("found_in") or []))
        ws.cell(row=r, column=6,
                value=rec.get("access") or rec.get("oa_status") or "unknown")
        ws.cell(row=r, column=7, value=rec.get("pmid") or "")
        for c in range(1, len(ALL_COLS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP if c in (1, 2) else TOP
    for col, w in zip("ABCDEFG", ALL_WIDTHS):
        ws.column_dimensions[col].width = w
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_COLS))}{ws.max_row}"
    style_header(ws, len(ALL_COLS))
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--selection", required=True, help="selection.json from Stage 5.1")
    ap.add_argument("--merged", required=True, help="merged.json from merge_results.py")
    ap.add_argument("--out", required=True, help="output .xlsx path")
    a = ap.parse_args()

    sel = json.loads(Path(a.selection).read_text(encoding="utf-8"))
    merged = json.loads(Path(a.merged).read_text(encoding="utf-8"))
    records = merged.get("records", merged) if isinstance(merged, dict) else merged

    wb = Workbook()
    sheet_selection(wb, sel)
    sheet_all(wb, records)

    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    items = sel.get("items", [])
    n_pdf = sum(1 for i in items if i.get("status") == STATUS_PDF)
    n_text = sum(1 for i in items if i.get("status") == STATUS_TEXT)
    n_miss = sum(1 for i in items if i.get("status") == STATUS_MISS)
    print(f"wrote {out}")
    print(f"  推荐文献: {len(items)} 行 "
          f"(PDF {n_pdf} / 全文非PDF {n_text} / 未下载 {n_miss})")
    print(f"  全部检索结果: {len(records)} 行")


if __name__ == "__main__":
    main()
